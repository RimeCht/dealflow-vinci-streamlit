import re
import unicodedata

from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


COLORS = {
    "ink": "243447",
    "muted": "5E6C7B",
    "line": "D9E2EA",
    "white": "FFFFFF",
    "algo1": "2367A5",
    "algo2": "00857C",
    "algo3": "6950A1",
    "pipeline": "3F4D5E",
    "manual": "B66A14",
    "green": "DDF3E4",
    "green_text": "17653A",
    "amber": "FFF0C7",
    "amber_text": "8A5600",
    "red": "FBE0E3",
    "red_text": "A4262C",
    "blue": "DEECF9",
    "blue_text": "174F7A",
    "purple": "E9E1F7",
    "purple_text": "55358C",
    "teal": "D7F0ED",
    "teal_text": "00665E",
    "gray": "E9EDF1",
    "gray_text": "4B5967",
    "zebra": "F7F9FB",
}

THIN_BORDER = Border(bottom=Side(style="thin", color=COLORS["line"]))


def normalize_text(value) -> str:
    text = str(value or "").strip()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[\s\-/]+", "_", text.upper())


def header_color(header: str) -> str:
    normalized = normalize_text(header).lower()
    if normalized.startswith("algo1_") or normalized in {"decision_finale"}:
        return COLORS["algo1"]
    if normalized.startswith("algo2_") or normalized.startswith("edda_"):
        return COLORS["algo2"]
    if normalized.startswith("algo3_"):
        return COLORS["algo3"]
    if normalized.startswith("pipeline_") or normalized in {
        "resultat_global",
        "etape_finale",
        "statut_pipeline",
        "prochaine_action",
    }:
        return COLORS["pipeline"]
    if normalized.startswith("manual_") or normalized.startswith("commentaire_"):
        return COLORS["manual"]
    return COLORS["ink"]


def tab_color(sheet_name: str) -> str:
    normalized = normalize_text(sheet_name)
    if "ALGO1" in normalized:
        return COLORS["algo1"]
    if "ALGO2" in normalized or "PREQUAL" in normalized:
        return COLORS["algo2"]
    if "ALGO3" in normalized or normalized in {
        "SEED",
        "CATALYST",
        "MATERIAUX",
        "AUTRE",
    }:
        return COLORS["algo3"]
    if "VERIFIER" in normalized:
        return "E5A100"
    if "ECARTER" in normalized or "NON_PREQUAL" in normalized:
        return "C83C45"
    return COLORS["pipeline"]


def status_style(value) -> tuple[str, str] | None:
    normalized = normalize_text(value)
    if not normalized:
        return None

    if "ORIENTEE_ALGO3" in normalized:
        if "CATALYST" in normalized:
            return COLORS["purple"], COLORS["purple_text"]
        if "MATERIAUX" in normalized:
            return COLORS["teal"], COLORS["teal_text"]
        if "SEED" in normalized:
            return COLORS["blue"], COLORS["blue_text"]
    if (
        normalized in {"A_GARDER", "PREQUALIFIEE", "ACTIVE", "FORT", "OUI"}
        or normalized.startswith("PREQUALIFIEE_")
        or normalized.startswith("TERMINE")
    ):
        return COLORS["green"], COLORS["green_text"]
    if normalized in {"SEED", "EARLY"}:
        return COLORS["blue"], COLORS["blue_text"]
    if normalized in {"CATALYST", "MATURE"}:
        return COLORS["purple"], COLORS["purple_text"]
    if normalized == "MATERIAUX":
        return COLORS["teal"], COLORS["teal_text"]
    if normalized in {"AUTRE", "INCONNU", "AUCUNE"}:
        return COLORS["gray"], COLORS["gray_text"]
    if "A_VERIFIER" in normalized or normalized in {"MOYEN", "PROBABLE", "INTERMEDIATE"}:
        return COLORS["amber"], COLORS["amber_text"]
    if (
        "A_ECARTER" in normalized
        or "ECARTEE" in normalized
        or "NON_PREQUALIFIEE" in normalized
        or "INACTIVE" in normalized
        or "RADIEE" in normalized
        or "ERREUR" in normalized
        or normalized in {"FAIBLE", "NON"}
    ):
        return COLORS["red"], COLORS["red_text"]
    return None


def is_status_column(header: str) -> bool:
    normalized = normalize_text(header).lower()
    return any(
        marker in normalized
        for marker in (
            "decision",
            "orientation",
            "resultat_global",
            "pipeline_status",
            "statut_pipeline",
            "edda_status",
            "strength",
            "maturity_stage",
            "materials_vertical",
            "passe_vers",
        )
    )


def is_long_text_column(header: str) -> bool:
    normalized = normalize_text(header).lower()
    return any(
        marker in normalized
        for marker in (
            "reason",
            "raison",
            "description",
            "comment",
            "missing_information",
            "infos_manquantes",
            "prochaine_action",
            "research_text",
            "summary",
            "matches",
        )
    )


def column_width(worksheet, column_index: int, header: str) -> float:
    normalized = normalize_text(header).lower()
    if is_long_text_column(header):
        return 48
    if any(marker in normalized for marker in ("startup", "company", "nom_", "site", "url")):
        return 27
    if any(marker in normalized for marker in ("source", "sector", "secteur", "matrix", "matrice")):
        return 25

    longest = len(str(header or ""))
    sample_end = min(worksheet.max_row, 80)
    for row_index in range(2, sample_end + 1):
        value = worksheet.cell(row=row_index, column=column_index).value
        if value is not None:
            longest = max(longest, min(len(str(value)), 24))
    return min(max(longest + 2, 12), 26)


def style_worksheet(worksheet, emphasize_rows: bool = False) -> None:
    if worksheet.max_column < 1:
        return

    worksheet.sheet_view.showGridLines = False
    worksheet.sheet_view.zoomScale = 85
    worksheet.freeze_panes = "C2" if emphasize_rows else "A2"
    worksheet.sheet_properties.tabColor = tab_color(worksheet.title)
    if worksheet.max_row >= 1:
        worksheet.auto_filter.ref = worksheet.dimensions

    headers = {}
    worksheet.row_dimensions[1].height = 32
    for column_index in range(1, worksheet.max_column + 1):
        cell = worksheet.cell(row=1, column=column_index)
        header = str(cell.value or "")
        headers[column_index] = header
        cell.fill = PatternFill("solid", fgColor=header_color(header))
        cell.font = Font(color=COLORS["white"], bold=True, size=10)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
        worksheet.column_dimensions[get_column_letter(column_index)].width = column_width(
            worksheet,
            column_index,
            header,
        )

    if emphasize_rows:
        for row_index in range(2, worksheet.max_row + 1):
            worksheet.row_dimensions[row_index].height = 30
            if row_index % 2 == 0:
                for column_index in range(1, worksheet.max_column + 1):
                    worksheet.cell(row=row_index, column=column_index).fill = PatternFill(
                        "solid",
                        fgColor=COLORS["zebra"],
                    )

    for column_index, header in headers.items():
        long_text = is_long_text_column(header)
        status_column = is_status_column(header)
        confidence_column = "confidence" in normalize_text(header).lower()
        link_column = any(
            marker in normalize_text(header).lower()
            for marker in ("site", "website", "url", "source_")
        )
        for row_index in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row=row_index, column=column_index)
            if long_text:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            if status_column:
                style = status_style(cell.value)
                if style:
                    fill_color, font_color = style
                    cell.fill = PatternFill("solid", fgColor=fill_color)
                    cell.font = Font(color=font_color, bold=True)
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if confidence_column and isinstance(cell.value, (int, float)):
                cell.number_format = "0%"
            if link_column and isinstance(cell.value, str) and cell.value.startswith(("http://", "https://")):
                cell.hyperlink = cell.value
                cell.font = Font(color=COLORS["algo1"], underline="single")

        if confidence_column and worksheet.max_row >= 2:
            letter = get_column_letter(column_index)
            worksheet.conditional_formatting.add(
                f"{letter}2:{letter}{worksheet.max_row}",
                ColorScaleRule(
                    start_type="num",
                    start_value=0,
                    start_color="F7C9CD",
                    mid_type="num",
                    mid_value=0.65,
                    mid_color="FFE4A3",
                    end_type="num",
                    end_value=1,
                    end_color="BDE5C8",
                ),
            )


def add_guide_sheet(workbook, title: str = "Résultats du dealflow") -> None:
    sheet_name = "Mode d'emploi"
    if sheet_name in workbook.sheetnames:
        del workbook[sheet_name]
    worksheet = workbook.create_sheet(sheet_name, 0)
    worksheet.sheet_view.showGridLines = False
    worksheet.sheet_properties.tabColor = COLORS["ink"]
    worksheet.freeze_panes = "A5"
    worksheet.merge_cells("A1:D1")
    worksheet["A1"] = title
    worksheet["A1"].fill = PatternFill("solid", fgColor=COLORS["ink"])
    worksheet["A1"].font = Font(color=COLORS["white"], bold=True, size=18)
    worksheet["A1"].alignment = Alignment(vertical="center")
    worksheet.row_dimensions[1].height = 42

    worksheet.merge_cells("A2:D2")
    worksheet["A2"] = (
        "Commencez par l'onglet Synthese equipe. Les autres onglets permettent "
        "de contrôler le détail et les preuves utilisées."
    )
    worksheet["A2"].font = Font(color=COLORS["muted"], italic=True, size=11)
    worksheet["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    worksheet.row_dimensions[2].height = 36

    headers = ["Étape", "Résultat", "Signification", "Action recommandée"]
    for index, value in enumerate(headers, 1):
        cell = worksheet.cell(row=4, column=index, value=value)
        cell.fill = PatternFill("solid", fgColor=COLORS["pipeline"])
        cell.font = Font(color=COLORS["white"], bold=True)
        cell.alignment = Alignment(vertical="center")

    rows = [
        ("Algo 1", "A_GARDER", "Éligible au premier filtre", "Passe à l'Algo 2"),
        ("Algo 1", "A_VERIFIER", "Éligibilité incertaine", "Vérification manuelle"),
        ("Algo 1", "A_ECARTER", "Hors critères initiaux", "Ne poursuit pas le pipeline"),
        ("Algo 2", "PREQUALIFIEE", "Alignement stratégique confirmé", "Passe à l'Algo 3"),
        ("Algo 2", "A_VERIFIER", "Une preuve importante manque", "Vérification manuelle"),
        ("Algo 2", "NON_PREQUALIFIEE", "Alignement insuffisant", "Ne poursuit pas le pipeline"),
        ("Algo 3", "Seed", "Startup jeune ou intermédiaire", "Étudier pour le programme Seed"),
        ("Algo 3", "Catalyst", "Startup mature et déployable", "Étudier pour Catalyst"),
        ("Algo 3", "Matériaux", "Innovation matériau prioritaire", "Orienter vers Matériaux"),
        ("Algo 3", "Autre", "Préqualifiée, autre orientation", "Orientation métier à définir"),
        ("Algo 3", "A_VERIFIER", "Maturité ou preuves insuffisantes", "Compléter les informations"),
    ]
    for row_index, values in enumerate(rows, 5):
        for column_index, value in enumerate(values, 1):
            cell = worksheet.cell(row=row_index, column=column_index, value=value)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        result_cell = worksheet.cell(row=row_index, column=2)
        style = status_style(result_cell.value)
        if style:
            fill_color, font_color = style
            result_cell.fill = PatternFill("solid", fgColor=fill_color)
            result_cell.font = Font(color=font_color, bold=True)
            result_cell.alignment = Alignment(horizontal="center", vertical="center")
        worksheet.row_dimensions[row_index].height = 28

    worksheet.column_dimensions["A"].width = 14
    worksheet.column_dimensions["B"].width = 24
    worksheet.column_dimensions["C"].width = 38
    worksheet.column_dimensions["D"].width = 38


def style_workbook(writer, include_guide: bool = True, title: str = "Résultats du dealflow") -> None:
    workbook = writer.book
    if include_guide:
        add_guide_sheet(workbook, title)

    for worksheet in workbook.worksheets:
        if worksheet.title == "Mode d'emploi":
            continue
        style_worksheet(
            worksheet,
            emphasize_rows=worksheet.title in {"Synthese equipe", "Synthèse équipe"},
        )
